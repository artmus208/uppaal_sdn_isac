from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"C:\Users\pEw\Documents\PHY уровень\ТинишовВС_Научный_проект_6G_ISAC_БПЛА.docx")


PROJECT_TITLE = (
    "Модели и методы интегрированного зондирования в 6G-сетях "
    "для обнаружения малоразмерных летательных аппаратов"
)

PARTICIPANT = "Тинишов Вадим Сергеевич"
ORG = "Национальный исследовательский университет ИТМО"
CATEGORY = "аспирант"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", sz="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, size=14, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=True):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if first_line:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, 14)
    return p


def add_bold_label_para(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.first_line_indent = Cm(0)
    r1 = p.add_run(label)
    set_run_font(r1, 14, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 14)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 16 if level == 1 else 14, bold=True, color="1F4D78" if level == 1 else "000000")
    return p


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(item)
        set_run_font(r, 14)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(item)
        set_run_font(r, 14)


def add_references(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        run = p.add_run(f"{i}. {item}")
        set_run_font(run, 13)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_margins(cell)
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 11, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, 10)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return table


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(1.5)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for name in ["Heading 1", "Heading 2", "Heading 3", "List Number", "List Bullet"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        if "Heading" in name:
            style.font.bold = True

    footer = sec.footer.paragraphs[0]
    footer.text = ""
    add_page_number(footer)


def build():
    doc = Document()
    configure_styles(doc)

    # Cover / theses page.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("МАТЕРИАЛЫ ДЛЯ ПУБЛИКАЦИИ (ТЕЗИСЫ)")
    set_run_font(r, 14, bold=True)

    add_bold_label_para(doc, "НАИМЕНОВАНИЕ НАУЧНОГО ПРОЕКТА: ", PROJECT_TITLE)
    add_bold_label_para(doc, "УЧАСТНИК КОНКУРСА: ", PARTICIPANT)
    add_bold_label_para(doc, "КАТЕГОРИЯ УЧАСТНИКА КОНКУРСА: ", CATEGORY)
    add_bold_label_para(doc, "ОРГАНИЗАЦИЯ: ", ORG)
    add_bold_label_para(
        doc,
        "ТЕЗИСЫ (АННОТАЦИЯ) НАУЧНОГО ПРОЕКТА: ",
        (
            "Проект посвящен исследованию и разработке моделей и методов интегрированного "
            "зондирования и связи в 6G-сетях для обнаружения малоразмерных летательных аппаратов, "
            "включая БПЛА с малой эффективной площадью рассеяния. Актуальность проекта определяется "
            "ростом числа сценариев применения беспилотных систем и одновременным увеличением рисков "
            "несанкционированного проникновения малоразмерных объектов в контролируемые зоны. "
            "Традиционные радиолокационные, оптические и акустические средства обнаружения требуют "
            "дополнительной инфраструктуры и не всегда обеспечивают устойчивую работу в урбанизированной "
            "среде, при многолучевом распространении, электромагнитных помехах и ограничениях по энергии. "
            "Концепция ISAC позволяет использовать инфраструктуру перспективных сетей 6G одновременно "
            "для передачи данных и радиозондирования среды. В рамках проекта предлагается системная "
            "модель 6G ISAC, включающая базовую станцию с массивной антенной решеткой, легальный "
            "зондирующий БПЛА и малоразмерный БПЛА как объект обнаружения. Научно-технический задел "
            "состоит в разработанной формальной иерархической модели SDN-управляемой ISAC-сети, где "
            "PHY-, MAC-, SDN/RIC- и сервисный уровни представлены аппроксимированными контрактными "
            "временными автоматами. Такая модель позволяет описывать вероятность обнаружения, ложные "
            "срабатывания, точность оценки, свежесть sensing-информации, состояние луча и канала как "
            "конечные классы и проверять своевременную реакцию системы на деградацию зондирования. "
            "Ожидаемый результат проекта — методика построения и проверки ISAC-подсистемы обнаружения "
            "малоразмерных летательных аппаратов с учетом ограничений спектра, энергии, вычислительных "
            "ресурсов и конфликтов между задачами связи и зондирования."
        ),
    )
    add_heading(doc, "1. Введение", 1)
    add_para(
        doc,
        (
            "Развитие сетей шестого поколения связано не только с ростом пропускной способности, "
            "снижением задержек и использованием высокочастотных диапазонов, но и с переходом к "
            "интеграции коммуникационных и сенсорных функций. Концепция Integrated Sensing and "
            "Communication (ISAC) предполагает, что радиосигналы сетевой инфраструктуры используются "
            "одновременно для передачи данных и получения информации о физической среде. Для задач "
            "обнаружения малоразмерных летательных аппаратов это направление имеет особое значение, "
            "поскольку инфраструктура 6G потенциально может выполнять функции распределенного "
            "радиозондирования без развертывания отдельной специализированной радиолокационной сети."
        ),
    )
    add_para(
        doc,
        (
            "Малоразмерные БПЛА являются сложными объектами обнаружения: они обладают малой эффективной "
            "площадью рассеяния, могут двигаться на малых высотах, скрываться в городской застройке, "
            "создавать доплеровские и угловые неоднозначности, а также попадать в зоны многолучевого "
            "распространения. При этом система обнаружения должна сохранять качество связи, не нарушать "
            "требования по задержке и пропускной способности, а также рационально использовать "
            "ограниченные спектральные, энергетические и вычислительные ресурсы."
        ),
    )
    add_para(
        doc,
        (
            "На практике возникает противоречие между точностью и устойчивостью радиозондирования, с "
            "одной стороны, и ресурсной эффективностью 6G-сети, с другой. Увеличение плотности пилотных "
            "или зондирующих сигналов, сужение луча и рост мощности могут повышать вероятность "
            "обнаружения цели, однако одновременно увеличивают накладные расходы, взаимные помехи и "
            "нагрузку на MAC- и SDN-уровни. Поэтому требуется не только радиофизическая модель канала, "
            "но и системная модель управления, позволяющая проверять своевременную реакцию сети на "
            "ухудшение sensing-KPI."
        ),
    )
    add_para(
        doc,
        (
            "Предлагаемый проект направлен на разработку таких моделей и методов. Его особенность "
            "состоит в соединении трех уровней описания: физического уровня ISAC, где оцениваются "
            "метрики обнаружения и состояния канала; уровня управления ресурсами, где распределяются "
            "временно-частотные и лучевые ресурсы между связью и зондированием; уровня формальной "
            "верификации, где деградации представляются конечными классами и проверяются как события "
            "контрактных временных автоматов."
        ),
    )

    add_heading(doc, "2. Основная часть", 1)
    add_heading(doc, "2.1. Степень разработанности темы", 2)
    add_para(
        doc,
        (
            "Современные исследования ISAC развиваются по нескольким направлениям. Первое направление "
            "связано с проектированием сигналов и фрейм-структур, позволяющих совмещать передачу "
            "полезной нагрузки и радиозондирование. Второе направление рассматривает sensing capacity, "
            "вероятность обнаружения, точность и свежесть sensing-информации как системные метрики. "
            "Третье направление связано с cooperative sensing, когда несколько узлов сети совместно "
            "формируют оценку цели. Четвертое направление исследует beam management для mmWave, "
            "sub-THz и THz-сценариев, где узкие лучи позволяют повысить разрешение, но создают риск "
            "потери сопровождения при мобильности объекта."
        ),
    )
    add_para(
        doc,
        (
            "Для темы обнаружения малоразмерных БПЛА указанные направления важны, но не закрывают "
            "полностью задачу системного проектирования. Радиофизическая модель может оценивать "
            "вероятность обнаружения и ложной тревоги, но сама по себе не показывает, как сеть должна "
            "перераспределять ресурсы при ухудшении качества зондирования. Методы управления лучом "
            "могут снижать накладные расходы, но без связи с MAC/SDN-уровнями трудно гарантировать "
            "ограниченное время реакции. Поэтому научная ниша проекта состоит в разработке моделей, "
            "которые связывают физические sensing-KPI, управление ресурсами и формальную проверку "
            "своевременной реакции 6G ISAC-сети."
        ),
    )
    add_table(
        doc,
        ["Направление", "Что уже исследуется", "Нерешенный вопрос для проекта"],
        [
            [
                "Signaling design для ISAC",
                "Совмещение пилотного, payload-assisted и joint sensing; выбор формы сигнала и структуры ресурса.",
                "Необходима связь параметров сигнала с конечными классами деградации и решениями MAC/SDN.",
            ],
            [
                "SensCAP и sensing-KPI",
                "Вероятность обнаружения, ложные тревоги, точность, coverage, sensing capacity и freshness.",
                "Требуется использовать эти KPI не только как радиометрики, но и как входы формальной модели реакции сети.",
            ],
            [
                "Cooperative sensing",
                "Совместное зондирование несколькими узлами и агрегация результатов.",
                "Нужна модель согласованности отчетов и давности информации при передаче к контроллеру.",
            ],
            [
                "Beam management",
                "Поиск, сопровождение и восстановление луча в высокочастотных диапазонах.",
                "Для малоразмерных БПЛА критично формализовать потерю луча, deadline восстановления и связь с обнаружением.",
            ],
            [
                "Timed automata и SDN-управление",
                "Формальная проверка дедлайнов, отсутствия тупиков, корректности переходов и реакций.",
                "Требуется адаптировать аппарат к ISAC-сценариям обнаружения, где деградация sensing влияет на сервис и ресурсы.",
            ],
        ],
        [2100, 3600, 3660],
    )

    add_heading(doc, "2.2. Цель, объект, предмет и задачи исследования", 2)
    add_para(
        doc,
        (
            "Цель исследования — разработать модели и методы интегрированного зондирования в сетях "
            "6G на основе концепции ISAC для обнаружения и сопровождения малоразмерных летательных "
            "аппаратов, обеспечивающие контролируемую вероятность обнаружения, ограничение ложных "
            "срабатываний, допустимую точность оценки и своевременную реакцию сети при ограниченных "
            "ресурсах спектра, энергии и вычислений."
        ),
    )
    add_para(
        doc,
        (
            "Объект исследования — инфраструктура перспективных 6G-сетей, включающая базовые станции "
            "с массивными антенными решетками, зондирующие БПЛА, каналы «воздух–земля» и "
            "«воздух–воздух», а также малоразмерные летательные аппараты как объекты обнаружения."
        ),
    )
    add_para(
        doc,
        (
            "Предмет исследования — модели каналов и эхо-сигналов, методы beam management, "
            "распределение ресурсов между communication и sensing, а также формальные модели "
            "своевременной реакции 6G ISAC-сети на деградацию качества зондирования."
        ),
    )
    add_para(doc, "Для достижения цели необходимо решить следующие задачи:", first_line=False)
    add_numbered(
        doc,
        [
            "Разработать системную модель 6G ISAC-сценария, включающую базовую станцию, легальный зондирующий БПЛА и малоразмерный БПЛА как объект обнаружения.",
            "Построить модели каналов «воздух–земля» и «воздух–воздух» с учетом mmWave/sub-THz диапазонов, многолучевости, доплеровского сдвига и помех.",
            "Определить набор sensing-KPI для обнаружения: вероятность обнаружения, вероятность ложной тревоги, точность оценки, CRB, давность sensing-информации и качество сопровождения луча.",
            "Разработать модель дискретизации непрерывных PHY-метрик в конечные классы для применения формальных методов проверки.",
            "Сформировать иерархическую модель ISAC-сети на основе аппроксимированных контрактных временных автоматов PHY, MAC, SDN/RIC и сервисного уровня.",
            "Разработать сценарии имитационного моделирования для оценки обнаружения малоразмерных БПЛА в условиях помех, блокировок, мобильности и ограниченных ресурсов.",
            "Сравнить предлагаемые методы с базовыми вариантами: статическим зондированием, раздельным проектированием связи и зондирования, а также реактивным восстановлением после деградации.",
            "Подготовить рекомендации по выбору параметров sensing, beam management и SDN/MAC-политик для построения практических 6G ISAC-систем обнаружения.",
        ],
    )

    add_heading(doc, "2.3. Научная новизна", 2)
    add_numbered(
        doc,
        [
            "Предлагается системная модель обнаружения малоразмерного БПЛА в 6G ISAC-сети, связывающая бистатическое радиозондирование, состояние канала, beam management и сетевые ограничения.",
            "Разрабатывается метод представления PHY-метрик обнаружения в виде конечных классов, пригодных для проверки временных контрактов: PdClass, RfaClass, AccClass, CRBClass, AoSClass, BeamClass и SensingState.",
            "Предлагается иерархическая модель ISAC-сети на основе контрактных временных автоматов, где своевременная реакция MAC/SDN на деградацию sensing проверяется формально.",
            "Обосновывается совместное рассмотрение точности обнаружения, свежести sensing-информации, стабильности луча и ресурсных ограничений как единой задачи управления 6G ISAC-системой.",
        ],
    )

    add_heading(doc, "3. Аналитические исследования", 1)
    add_para(
        doc,
        (
            "Аналитическая часть проекта строится вокруг модели, в которой малоразмерный БПЛА "
            "обнаруживается с использованием сигналов 6G ISAC. Базовая станция и зондирующий БПЛА "
            "формируют распределенную sensing-конфигурацию, а контроллер сети получает не непрерывные "
            "радиофизические величины, а конечные классы состояния. Такой подход позволяет отделить "
            "радиофизическую оценку от логики управления и формальной проверки."
        ),
    )
    add_para(
        doc,
        (
            "В качестве исходных метрик рассматриваются SINR канала связи, sensing-SINR, вероятность "
            "обнаружения Pd, вероятность ложной тревоги Rfa, точность оценки координат и скорости, "
            "CRB, вероятность потери луча, накладные расходы beam management и Age of Sensing. "
            "Каждая из этих величин после оценки внешним estimator-слоем переводится в конечный класс. "
            "Например, вероятность обнаружения может принимать классы OK, LOW или FAILED, а давность "
            "sensing-информации — FRESH, STALE или EXPIRED."
        ),
    )
    add_table(
        doc,
        ["Компонент модели", "Содержание", "Связь с обнаружением БПЛА"],
        [
            [
                "A_CH",
                "Автомат состояния канала: nominal, interference-limited, multipath-limited, blockage, outage.",
                "Позволяет учитывать помехи, блокировки и многолучевость, влияющие на обнаружение малой цели.",
            ],
            [
                "A_SIG",
                "Автомат сигнальной конфигурации: pilot-based, payload-assisted, reconfiguring, limited.",
                "Описывает режим использования радиосигнала для связи и зондирования.",
            ],
            [
                "A_BM",
                "Автомат beam management: search, track, lock, predict, recovering, failed.",
                "Формализует потерю и восстановление луча при движении БПЛА.",
            ],
            [
                "A_SQ",
                "Автомат качества зондирования: ProbabilityLimited, FalseAlarmLimited, AccuracyLimited, FreshnessLimited.",
                "Связывает sensing-KPI с событиями деградации обнаружения.",
            ],
            [
                "A_PH",
                "Агрегирующий автомат PHY-состояния.",
                "Формирует единый отчет для MAC/SDN о состоянии канала, луча и sensing-функции.",
            ],
        ],
        [1800, 3650, 3910],
    )
    add_para(
        doc,
        (
            "На более высоких уровнях модель расширяется автоматами MAC/resource scheduling, SDN/RIC "
            "и application/service layer. MAC-уровень принимает локальные решения о распределении "
            "радиоресурсов между communication и sensing. SDN/RIC выбирает политику реакции: усиление "
            "зондирования, ограниченный режим, перепланирование ресурсов или отказ с явной причиной. "
            "Сервисный уровень задает требования к минимальной вероятности обнаружения, допустимой "
            "ложной тревоге, свежести sensing-информации и задержке реакции."
        ),
    )

    add_heading(doc, "4. Экспериментальные исследования", 1)
    add_heading(doc, "4.1. Условия и постановка эксперимента", 2)
    add_para(
        doc,
        (
            "Экспериментальная проверка предполагается в форме аналитико-имитационного моделирования. "
            "В качестве базового сценария рассматривается городская или промышленная зона, где "
            "базовая станция 6G и легальный зондирующий БПЛА должны обнаружить малоразмерный "
            "нелегальный БПЛА. Варьируются дальность, скорость цели, уровень помех, наличие блокировок, "
            "многолучевость, ширина луча, период sensing-обновлений и доля ресурса, выделенная на "
            "радиозондирование."
        ),
    )
    add_para(doc, "Для сравнения целесообразно использовать следующие базовые стратегии:", first_line=False)
    add_bullets(
        doc,
        [
            "S1 — раздельное проектирование связи и зондирования без SDN/MAC-согласования.",
            "S2 — статическое распределение sensing-ресурса без реакции на деградацию.",
            "S3 — реактивное восстановление после потери луча или деградации Pd/Rfa.",
            "S4 — предлагаемая контрактная ISAC-модель с ограниченными дедлайнами реакции MAC/SDN.",
        ],
    )
    add_heading(doc, "4.2. Метрики оценки", 2)
    add_table(
        doc,
        ["Метрика", "Смысл", "Ожидаемое использование"],
        [
            ["Pd", "Вероятность правильного обнаружения цели.", "Основная метрика качества обнаружения БПЛА."],
            ["Rfa", "Вероятность ложной тревоги.", "Контроль ложных срабатываний в городской среде."],
            ["Acc", "Точность оценки координат и скорости.", "Оценка пригодности модели для сопровождения цели."],
            ["CRB", "Теоретическая нижняя граница ошибки оценки.", "Сравнение достижимой точности при разных режимах сигнала."],
            ["AoS", "Возраст sensing-информации.", "Контроль свежести данных для SDN/RIC и сервисного уровня."],
            ["p_mis", "Риск рассогласования луча.", "Оценка устойчивости beam management при движении БПЛА."],
            ["D_react", "Время реакции MAC/SDN на деградацию sensing.", "Проверка bounded response в контрактной автоматной модели."],
            ["ResourceShare", "Доля ресурса, выделенная на sensing.", "Оценка компромисса между связью и зондированием."],
        ],
        [1700, 3700, 3960],
    )
    add_heading(doc, "4.3. Ожидаемые результаты проверки", 2)
    add_para(
        doc,
        (
            "Экспериментальная часть должна показать, при каких параметрах канала, луча и распределения "
            "ресурса система сохраняет допустимые классы Pd, Rfa, Acc и AoS. Отдельно будет проверяться, "
            "что при переходе sensing quality в предупреждающий или критический класс MAC/SDN-уровни "
            "не остаются в неопределенном состоянии, а выбирают один из разрешенных исходов: sensing boost, "
            "constrained mode, reconfiguration или явный отказ политики."
        ),
    )
    add_para(
        doc,
        (
            "В отличие от чисто радиофизического моделирования, предлагаемая проверка оценивает не только "
            "значения метрик обнаружения, но и управляемость системы: отсутствие тупиков, соблюдение "
            "дедлайнов, корректность отчетов PHY и MAC, а также согласованность решений SDN/RIC с "
            "требованиями сервисного уровня."
        ),
    )

    add_heading(doc, "5. Информация об опыте и научно-техническом заделе", 1)
    add_para(
        doc,
        (
            "По теме проекта сформирован научно-технический задел в виде формальной иерархической модели "
            "SDN-управляемой 6G ISAC-сети. Наиболее проработанным элементом является PHY-уровень, где "
            "описаны автоматы состояния канала, сигнальной конфигурации, beam management, качества "
            "радиозондирования и агрегированного PHY-состояния. Непрерывные радиофизические величины "
            "не подменяются логикой автоматов: они предварительно оцениваются внешним измерительным или "
            "аналитическим слоем, а затем отображаются в конечные классы."
        ),
    )
    add_para(
        doc,
        (
            "Сформированы классы, непосредственно связанные с задачей обнаружения малоразмерных летательных "
            "аппаратов: PdClass, RfaClass, AccClass, CRBClass, AoSClass, BeamClass, SensingState и PHYState. "
            "Это позволяет описывать ухудшение вероятности обнаружения, рост ложных тревог, потерю точности, "
            "устаревание sensing-информации и потерю луча как проверяемые события системы."
        ),
    )
    add_para(
        doc,
        (
            "Также разработан черновой вариант уровней MAC/resource scheduling, SDN/RIC и application/service "
            "layer. В модели задана композиция A_TOTAL = A_SVC || A_SDN || A_MAC || A_PHY || A_ENV, где "
            "A_ENV является общим автоматом внешней среды с локальными проекциями на уровни. Такая структура "
            "предотвращает несогласованность, при которой разные уровни сети видят разные версии одного и "
            "того же физического события."
        ),
    )
    add_para(
        doc,
        (
            "Апробация постановки диссертационной темы частично выполнена в рамках обсуждения положений "
            "на LV научной и учебно-методической конференции Университета ИТМО. Публикации по теме "
            "диссертационного исследования находятся в подготовке. Поэтому заявляемый на данном этапе "
            "результат следует формулировать как разработку формальной модели и методики проверки ISAC-"
            "подсистемы, а не как завершенную экспериментально подтвержденную систему обнаружения."
        ),
    )

    add_heading(doc, "6. Выводы и ожидаемая практическая значимость", 1)
    add_para(
        doc,
        (
            "В рамках проекта предлагается перейти от разрозненного рассмотрения радиоканала, "
            "beam management и сетевого управления к единой модели 6G ISAC-системы обнаружения "
            "малоразмерных летательных аппаратов. Основной научный результат должен состоять в "
            "формализации связи между sensing-KPI, состоянием канала, управлением лучом и решениями "
            "MAC/SDN-уровней."
        ),
    )
    add_para(
        doc,
        (
            "Практическая значимость проекта заключается в возможности использовать результаты при "
            "проектировании систем мониторинга воздушного пространства на базе будущей 6G-инфраструктуры, "
            "в том числе для умных городов, промышленных зон, логистических коридоров и объектов с "
            "повышенными требованиями к контролю воздушной обстановки. Для Санкт-Петербурга, как крупного "
            "научно-образовательного, промышленного и телекоммуникационного центра, такая тематика связана "
            "с развитием отечественных решений в области перспективных сетей связи, радиомониторинга и "
            "безопасности городской инфраструктуры."
        ),
    )
    add_para(
        doc,
        (
            "Ожидаемый итог проекта — набор моделей, алгоритмических процедур и сценариев проверки, "
            "позволяющих выбирать параметры интегрированного зондирования, оценивать компромисс между "
            "связью и обнаружением, а также формально проверять, что система своевременно реагирует на "
            "деградацию sensing-функции."
        ),
    )

    add_heading(doc, "Список используемой литературы", 1)
    refs = [
        "Li Y., Zhang Y., Masouros C., Pollin S., Liu F. Rethinking Signaling Design for ISAC: From Pilot-Based to Payload-Based Sensing // IEEE Communications Standards Magazine. In press. 2026.",
        "Liu G. et al. SensCAP: A Systematic Sensing Capability Performance Metric for 6G ISAC // IEEE Internet of Things Journal. 2024. Vol. 11, No. 18. P. 29438–29454.",
        "Liu G. et al. Cooperative Sensing for 6G ISAC: Concept, Key Technologies, Performance Evaluation, and Field Trial // Engineering. 2026. Vol. 56. P. 130–148.",
        "Huang Y., Xu J., Badiu M. A., Chen G., Coon J., Alouini M.-S. ISAC-Enabled Low-Overhead Beam Management: Performance Analysis and Pilot Optimization // IEEE Transactions on Wireless Communications. 2026. Vol. 25. P. 10702–10715.",
        "Alur R., Dill D. L. A Theory of Timed Automata // Theoretical Computer Science. 1994. Vol. 126, No. 2. P. 183–235.",
        "UPPAAL Documentation. Query Semantics: Symbolic Queries. Accessed 2026. URL: https://docs.uppaal.org/language-reference/query-semantics/symb_queries/.",
        "Tigane T., Sadovykh A., Bensalem S., Bozga M. Dynamic Timed Automata for Reconfigurable System Modeling. 2023.",
        "Mosudi O. A., Popoola S. I., Atayero A. A., Elngar A. A. SDN for 5G and Beyond: Transforming Network Architecture. 2025.",
        "Anand M., Vestal S. Formal Modeling and Analysis of the AFDX Frame Management Design // Proceedings of the IEEE International Conference on Engineering of Complex Computer Systems. 2007.",
        "Тинишов В. С. Текст диссертации: «Исследование и разработка моделей и методов интегрированного зондирования в 6G-сетях для обнаружения малоразмерных летательных аппаратов». Рабочий фрагмент. 2026.",
    ]
    add_references(doc, refs)

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
