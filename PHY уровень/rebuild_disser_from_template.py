from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FRONT_MATTER_PARAGRAPHS = 19


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def clear_after_front_matter(doc: Document):
    for p in list(doc.paragraphs)[FRONT_MATTER_PARAGRAPHS:]:
        remove_paragraph(p)


def add_plain_paragraph(doc: Document, text: str, align=WD_ALIGN_PARAGRAPH.LEFT, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=doc.styles["Normal"])
    p.alignment = align
    pf = p.paragraph_format
    pf.left_indent = None
    pf.right_indent = None
    pf.first_line_indent = None
    pf.space_before = None
    pf.space_after = None
    pf.line_spacing = None

    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        rest = text[len(bold_prefix):]
        if rest:
            p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_center_heading(doc: Document, text: str, bold=True):
    p = add_plain_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        run.bold = bold
    return p


def add_left_heading(doc: Document, text: str):
    p = add_plain_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT)
    for run in p.runs:
        run.bold = True
    return p


def normalize_line(line: str) -> str:
    return line.strip().replace("**", "").replace("`", "")


def ru(hex_text: str) -> str:
    return bytes.fromhex(hex_text).decode("utf-8")


def make_architecture_image(path: Path):
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1200, 860
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font_title = ImageFont.truetype(r"C:\Windows\Fonts\timesbd.ttf", 38)
        font_box = ImageFont.truetype(r"C:\Windows\Fonts\times.ttf", 31)
        font_small = ImageFont.truetype(r"C:\Windows\Fonts\times.ttf", 23)
    except Exception:
        font_title = font_box = font_small = ImageFont.load_default()

    draw.text((width // 2, 48), "Hierarchical 6G-ISAC architecture", font=font_title, fill="black", anchor="ma")
    boxes = [
        ("Application / SLA", "UAV-detection service", 145),
        ("SDN / RIC control plane", "policy, recovery, rollback", 285),
        ("MAC / resource scheduling", "queues, allocation, sensing boost", 425),
        ("PHY abstraction layer", "A_CH, A_SIG, A_BM, A_SQ, A_PH", 565),
        ("Physical environment and external PHY estimators", "channel, beam, Pd, Rfa, CRB, AoS", 705),
    ]
    box_w, box_h = 930, 95
    x0 = (width - box_w) // 2
    for index, (main, sub, y) in enumerate(boxes):
        fill = (235, 242, 251) if index % 2 == 0 else (242, 247, 242)
        draw.rounded_rectangle([x0, y, x0 + box_w, y + box_h], radius=18, fill=fill, outline=(55, 85, 120), width=3)
        draw.text((width // 2, y + 30), main, font=font_box, fill="black", anchor="ma")
        draw.text((width // 2, y + 65), sub, font=font_small, fill=(70, 70, 70), anchor="ma")
        if index < len(boxes) - 1:
            x = width // 2
            y1 = y + box_h + 8
            y2 = boxes[index + 1][2] - 8
            draw.line([x, y1, x, y2], fill=(40, 40, 40), width=4)
            draw.polygon([(x - 12, y2 - 16), (x + 12, y2 - 16), (x, y2)], fill=(40, 40, 40))

    draw.text(
        (width // 2, height - 35),
        "Reports propagate upward; SLA-significant decisions are made by upper layers",
        font=font_small,
        fill=(70, 70, 70),
        anchor="ma",
    )
    image.save(path)


def add_picture_with_caption(doc: Document, image_path: Path, caption: str):
    p = doc.add_paragraph(style=doc.styles["Normal"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = None
    p.paragraph_format.space_after = None
    p.add_run().add_picture(str(image_path), width=Cm(14.8))
    add_plain_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_math_omml(paragraph, text: str):
    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    omath.append(mr)
    paragraph._p.append(omath)


def clear_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tc_pr.append(borders)


def add_formula(doc: Document, expression: str, number: str):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(14.0)
    table.columns[1].width = Cm(2.0)
    left, right = table.cell(0, 0), table.cell(0, 1)
    clear_cell_borders(left)
    clear_cell_borders(right)

    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_before = None
    lp.paragraph_format.space_after = None
    add_math_omml(lp, expression)

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = None
    rp.paragraph_format.space_after = None
    run = rp.add_run(f"({number})")
    run.font.size = Pt(14)


def build(template_path: Path, md_path: Path, out_path: Path):
    doc = Document(template_path)
    clear_after_front_matter(doc)

    figure_path = out_path.with_name("isac_architecture_figure.png")
    make_architecture_image(figure_path)

    raw_lines = md_path.read_text(encoding="utf-8").splitlines()
    started = False
    pending_general_heading: str | None = None
    in_code = False
    code_lines: list[str] = []
    figure_counter = 0
    current_section = "0.0"
    formula_counts: dict[str, int] = {}

    for raw in raw_lines:
        raw_stripped = raw.strip()
        if raw_stripped.startswith("```"):
            if in_code:
                block = [normalize_line(c) for c in code_lines if normalize_line(c)]
                if block[:1] == ["Application / SLA"]:
                    figure_counter += 1
                    intro = ru("d090d180d185d0b8d182d0b5d0bad182d183d180d0b020d180d0b0d181d181d0bcd0b0d182d180d0b8d0b2d0b0d0b5d0bcd0bed0b920d181d0b8d181d182d0b5d0bcd18b20d0bfd180d0b5d0b4d181d182d0b0d0b2d0bbd0b5d0bdd0b020d0bdd0b020d180d0b8d181d183d0bdd0bad0b520")
                    add_plain_paragraph(doc, f"{intro}{figure_counter}.")
                    cap_a = ru("d0a0d0b8d181d183d0bdd0bed0ba20")
                    cap_b = ru("202d20d098d0b5d180d0b0d180d185d0b8d187d0b5d181d0bad0b0d18f20d0b0d180d185d0b8d182d0b5d0bad182d183d180d0b02036472d4953414320d181d0b8d181d182d0b5d0bcd18b20d0bed0b1d0bdd0b0d180d183d0b6d0b5d0bdd0b8d18f20d0bcd0b0d0bbd0bed180d0b0d0b7d0bcd0b5d180d0bdd18bd18520d091d09fd09bd090")
                    add_picture_with_caption(
                        doc,
                        figure_path,
                        f"{cap_a}{figure_counter}{cap_b}",
                    )
                else:
                    for expr in block:
                        formula_counts[current_section] = formula_counts.get(current_section, 0) + 1
                        add_formula(doc, expr, f"{current_section}.{formula_counts[current_section]}")
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(raw.rstrip())
            continue

        line = normalize_line(raw)
        if not line:
            continue

        if not started:
            if line.startswith("# "):
                continue
            if line.startswith("## "):
                started = True
                continue
            continue

        if line.startswith("### "):
            pending_general_heading = line[4:].strip()
            continue

        if line.startswith("# "):
            add_center_heading(doc, line[2:].strip(), bold=True)
            pending_general_heading = None
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            add_left_heading(doc, heading)
            m = re.match(r"(\d+)\.(\d+)\.", heading)
            if m:
                current_section = f"{m.group(1)}.{m.group(2)}"
            pending_general_heading = None
            continue

        if pending_general_heading:
            prefix = pending_general_heading
            if not prefix.endswith("."):
                prefix += "."
            add_plain_paragraph(doc, f"{prefix} {line}", bold_prefix=prefix)
            pending_general_heading = None
        else:
            add_plain_paragraph(doc, line)

    doc.save(out_path)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
    print(sys.argv[3])
