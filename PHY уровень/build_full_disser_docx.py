from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def set_run(run, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = bold
    run.italic = italic


def add_runs_from_inline_md(paragraph, text: str, *, bold_all=False):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run(run, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run(run, bold=bold_all)


def paragraph(doc: Document, text: str, *, align=None, bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = None
    add_runs_from_inline_md(p, text, bold_all=bold)
    return p


def build(md_path: Path, out_path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = None
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_buffer: list[str] = []

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("```"):
            if in_code:
                for code_line in code_buffer:
                    paragraph(doc, code_line, align=WD_ALIGN_PARAGRAPH.LEFT)
                code_buffer.clear()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buffer.append(raw.rstrip())
            continue

        if line.startswith("# "):
            paragraph(doc, line[2:].strip(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=6, space_after=6)
        elif line.startswith("## "):
            paragraph(doc, line[3:].strip(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=6, space_after=6)
        elif line.startswith("### "):
            paragraph(doc, line[4:].strip(), align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, space_before=6, space_after=3)
        elif re.match(r"^\d+\.\s+", line):
            paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("- "):
            paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.save(out_path)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
    print(sys.argv[2])
