from pathlib import Path
import subprocess,os,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
P=Path(__file__).resolve().parent.parent
pandoc=P/'build/tools/pypandoc/files/pandoc.exe'
if not pandoc.exists():
    import pypandoc
    pandoc=Path(os.environ.get('PANDOC_PATH') or shutil.which('pandoc') or pypandoc.get_pandoc_path())
out=P/'reviewer-response.docx'
subprocess.run([str(pandoc),str(P/'reviewer-response.md'),'-o',str(out)],check=True)
doc=Document(out)
for section in doc.sections:
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.8);section.bottom_margin=Inches(.8)
    section.left_margin=Inches(.9);section.right_margin=Inches(.9)
    section.header_distance=Inches(.35);section.footer_distance=Inches(.35)
for style in doc.styles:
    if style.type in [1,2]:
        style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
for name in ['Normal','Body Text','First Paragraph','Compact']:
    s=doc.styles[name];s.font.size=Pt(11)
    s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.04
    s.paragraph_format.widow_control=True
for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.keep_with_next=True
for p in doc.paragraphs:
    if p.text.startswith('Status:'):p.paragraph_format.keep_with_next=True
for table in doc.tables:
    table.autofit=False
    for ci,w in enumerate([2.1,4.6]):table.columns[ci].width=Inches(w)
    for ri,row in enumerate(table.rows):
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches([2.1,4.6][ci]);pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            if ri==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E6EAF0');pr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(10);run.bold=(ri==0)
for s in doc.sections:
    head=s.header.paragraphs[0];head.text='MoNeTec 2026 | Manuscript 18100004273'
    head.style=doc.styles['Caption']
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('Response to reviewers | ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='Response to Reviewers — Manuscript 18100004273'
doc.core_properties.author='Artur Mustafin; Vadim Tinishov'
doc.save(out)
print(out)
