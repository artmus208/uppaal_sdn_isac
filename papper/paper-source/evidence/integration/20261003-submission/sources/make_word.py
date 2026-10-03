"""Export the built TeX revision, resolved refs and exact BibTeX bibliography to Word."""
from pathlib import Path
import re,subprocess,json,hashlib,shutil,os
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
P=Path(__file__).resolve().parent.parent
R=P.parents[2]
pandoc=P/'build/tools/pypandoc/files/pandoc.exe'
if not pandoc.exists():
    import pypandoc
    pandoc=Path(os.environ.get('PANDOC_PATH') or shutil.which('pandoc') or pypandoc.get_pandoc_path())
tex=(R/'levels_tex/samplepaper.tex').read_text(encoding='utf-8')
aux=(P/'build/samplepaper.aux').read_text(encoding='utf-8')
refs=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',aux))
cites=dict(re.findall(r'\\bibcite\{([^}]+)\}\{([^}]+)\}',aux))
tex=re.sub(r'\\(?:eqref|ref)\{([^}]+)\}',lambda m:refs[m[1]],tex)
tex=re.sub(r'\\cite\{([^}]+)\}',lambda m:'['+', '.join(cites[k.strip()] for k in m[1].split(','))+']',tex)
tex=re.sub(r'\\titlerunning\{[^}]*\}|\\authorrunning\{[^}]*\}', '',tex)
tex=re.sub(r'\\author\{.*?\\institute\{.*?\\maketitle',r'\\author{Artur Mustafin and Vadim Tinishov}\n\\maketitle\n\\begin{center}\nITMO University, Faculty of Software Engineering and Computer Systems, Saint Petersburg, Russia\\\\\n\\texttt{armustafin@itmo.ru}; \\texttt{vadimtdot@gmail.com}\\\\\nORCID: 0009-0008-8322-0864; 0009-0003-7319-1944\n\\end{center}',tex,flags=re.S)
tex=tex.replace(r'\begin{abstract}',r'\section*{Abstract}').replace(r'\end{abstract}','')
tex=re.sub(r'\\keywords\{([^}]+)\}',lambda m:r'\paragraph{Keywords} '+m[1].replace(r'\and',';'),tex)
tex=tex.replace(r'\begin{credits}','').replace(r'\end{credits}','').replace(r'\discintname','Disclosure of Interests')
tex=re.sub(r'\\bibliographystyle\{[^}]+\}', '',tex)
bbl=(P/'build/samplepaper.bbl').read_text(encoding='utf-8')
bbl=re.sub(r'\\begin\{thebibliography\}\{[^}]*\}',r'\\section*{References}\n\\begin{enumerate}',bbl)
bbl=re.sub(r'\\bibitem\{[^}]+\}',r'\\item ',bbl).replace(r'\end{thebibliography}',r'\end{enumerate}')
tex=re.sub(r'\\bibliography\{[^}]+\}',lambda _:bbl,tex)
for stem in ['phy','mac','sdn','app']:
    tex=tex.replace(r'\figpath '+stem+'.pdf',(P/'figures'/f'{stem}.png').as_posix())
# Pandoc understands tabular, p-columns and booktabs. Give explicit widths in
# place of tabularx's stretchy X column; the actual content is unchanged.
tex=re.sub(r'\\begin\{tabularx\}\{\\textwidth\}\{(@\{\}[^\n]*?@\{\})\}',lambda m:r'\begin{tabular}{'+m[1].replace('X','p{5.5cm}')+'}',tex)
tex=tex.replace(r'\end{tabularx}',r'\end{tabular}')
conversion=P/'build/word-source.tex';conversion.write_text(tex,encoding='utf-8')
astpath=P/'build/word-source.json'
subprocess.run([str(pandoc),str(conversion),'-f','latex','-t','json','-o',str(astpath)],check=True)
ast=json.loads(astpath.read_text(encoding='utf-8'))
# Normalize the mathematical syntax through Pandoc's own OMML path. Explicit
# mixed interval delimiters and text A[] avoid LibreOffice's formula parser errors.
def normalize_math(o):
    if isinstance(o,dict):
        if o.get('t')=='Math':
            s=o['c'][1]
            s=re.sub(r'\\(?:begin|end)\{equation\}|\\label\{[^}]+\}','',s).strip()
            s=s.replace('A[]',r'\text{A[]}')
            if re.fullmatch(r'[\[(].+[\])]',s):s=r'\left'+s[:-1]+r'\right'+s[-1]
            o['c'][1]=s
        for v in o.values():normalize_math(v)
    elif isinstance(o,list):
        for v in o:normalize_math(v)
normalize_math(ast)
blocks=[]
for b in ast['blocks']:
    if b.get('t')!='Para':
        blocks.append(b);continue
    pending=[]
    def flush():
        while pending and pending[0]['t'] in ['Space','SoftBreak']:pending.pop(0)
        while pending and pending[-1]['t'] in ['Space','SoftBreak']:pending.pop()
        if pending:blocks.append({'t':'Para','c':list(pending)})
        pending.clear()
    for inline in b['c']:
        if inline.get('t')=='Math' and inline['c'][0]['t']=='DisplayMath':
            flush();s=inline['c'][1]
            if r'\begin{align}' in s:
                s=s.replace(r'\begin{align}','').replace(r'\end{align}','')
                for row in s.split(r'\\'):
                    row=re.sub(r'\\label\{[^}]*\}','',row).strip().lstrip('&')
                    blocks.append({'t':'Para','c':[{'t':'Math','c':[{'t':'DisplayMath'},row]}]})
            else:blocks.append({'t':'Para','c':[inline]})
        else:pending.append(inline)
    flush()
ast['blocks']=blocks
section=0
for b in ast['blocks']:
    if b.get('t')=='Header' and b['c'][0]==1 and 'unnumbered' not in b['c'][1][1]:
        section+=1;b['c'][2].insert(0,{'t':'Str','c':str(section)+'  '})
astpath.write_text(json.dumps(ast),encoding='utf-8')
def count_kind(obj,kind):
    if isinstance(obj,dict):return (obj.get('t')==kind)+sum(count_kind(v,kind) for v in obj.values())
    if isinstance(obj,list):return sum(count_kind(v,kind) for v in obj)
    return 0
assert count_kind(ast,'Table')==6, count_kind(ast,'Table')
assert count_kind(ast,'Image')==4, count_kind(ast,'Image')
out=P/'MoNoTeC-2026-paper.docx'
subprocess.run([str(pandoc),str(astpath),'-f','json','-t','docx','-o',str(out)],check=True)
doc=Document(out)
# Match the proceedings text block: 122 x 193 mm, 10 pt body.
for s in doc.sections:
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(4.4);s.right_margin=Cm(4.4)
    s.top_margin=Cm(4.6);s.bottom_margin=Cm(5.8)
    s.header_distance=Cm(3.5);s.footer_distance=Cm(4.6)
for style in doc.styles:
    if style.type in (1,2):
        style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
for name in ['Normal','Body Text','First Paragraph','Compact']:
    st=doc.styles[name];st.font.size=Pt(10)
    st.paragraph_format.line_spacing=1.0;st.paragraph_format.space_after=Pt(1)
    st.paragraph_format.widow_control=True
for name,size in [('Title',14),('Heading 1',12),('Heading 2',11),('Heading 3',10)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['Caption','Image Caption','Table Caption']:
    if name in doc.styles:
        doc.styles[name].font.size=Pt(9)
        doc.styles[name].paragraph_format.space_before=Pt(4)
        doc.styles[name].paragraph_format.space_after=Pt(6)
for p in doc.paragraphs:
    if p.style.name in ['Body Text','First Paragraph','Normal']:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    if p.style.name=='Table Caption':p.paragraph_format.keep_with_next=True
    if p.style.name in ['Image Caption','Caption']:p.paragraph_format.keep_with_next=False
fig_no=table_no=0
for p in doc.paragraphs:
    if p.style.name=='Image Caption':
        fig_no+=1;prefix=f'Fig. {fig_no}. '
    elif p.style.name=='Table Caption':
        table_no+=1;prefix=f'Table {table_no}. '
    else:continue
    r=p.add_run(prefix);p._p.insert(1,r._r);r.bold=True
figure_alts=[p.text for p in doc.paragraphs if p.style.name=='Image Caption']
for fi,shape in enumerate(doc.inline_shapes):
    ratio=shape.height/shape.width;shape.width=Cm(12.2);shape.height=int(shape.width*ratio)
    shape._inline.docPr.set('descr',figure_alts[fi])
for p in doc.paragraphs:
    if p.text.startswith('ITMO University'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:r.font.size=Pt(9)
# Resolved equation numbers remain editable text next to native OMML.
eq=0
for p in doc.paragraphs:
    if p._p.xpath('./m:oMathPara'):
        math=p._p.xpath('./m:oMathPara')[0]
        numbers=f'({eq+1})'
        p.add_run('  '+numbers)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;eq+=1
for ti,table in enumerate(doc.tables):
    table.autofit=False
    widths=[[3.1,1.7,7.4],[2.1,3.4,1.8,4.9],[3.7,1.2,1.2,6.1],[2.8,4.1,5.3],[.8,1.6,1.5,1.5,1.7,5.1],[1.2,5.0,6.0]][ti]
    for rowi,row in enumerate(table.rows):
        trPr=row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if rowi==0:trPr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            pad='30' if ti==5 else '55'
            for edge,val in [('top',pad),('bottom',pad),('left','65'),('right','65')]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if rowi==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EAF0');pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size=Pt(9);run.bold=(rowi==0)
    for ci,w in enumerate(widths):table.columns[ci].width=Cm(w)
in_refs=False
for p in doc.paragraphs:
    if p.text=='References':in_refs=True
    if in_refs:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
for s in doc.sections:
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='A Hierarchical Timed-Automata Model for SDN-Managed Resource Orchestration in 6G ISAC Networks'
doc.core_properties.author='Artur Mustafin; Vadim Tinishov'
doc.save(out)
stats={'source_sha256':hashlib.sha256((R/'levels_tex/samplepaper.tex').read_bytes()).hexdigest(),'tables':len(doc.tables),'images':len(doc.inline_shapes),'omml_nodes':len(doc._element.xpath('.//m:oMath')),'display_equation_blocks':eq,'citations_resolved':cites,'refs_resolved':refs}
(P/'checks/word-conversion.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
print(json.dumps(stats))
