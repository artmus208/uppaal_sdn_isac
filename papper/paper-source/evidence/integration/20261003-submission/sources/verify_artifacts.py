from pathlib import Path
import sys,json,hashlib,re,zipfile
P=Path(__file__).resolve().parent.parent;R=P.parents[2]
sys.path.insert(0,str(P/'build/tools'))
import pymupdf
from docx import Document
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
source=R/'levels_tex/samplepaper.tex'
stats=json.loads((P/'checks/word-conversion.json').read_text())
assert stats['source_sha256']==sha(source)
log=(P/'checks/latex-build.log').read_text(encoding='utf-8')
assert 'Overfull' not in log
assert not re.search(r'(Citation|Reference) .+ undefined',log)
tex=source.read_text(encoding='utf-8')
assert len(re.findall(r'\\begin\{table\}',tex))==6
assert len(re.findall(r'\\begin\{figure\}',tex))==4
assert len(stats['citations_resolved'])==9
models=[R/'evidence/instantiation/uav-service-completion-candidate/model.xml',P/'sources/layout-model.xml']
for path in models:assert sha(path)=='b8ab50e112b71491f8242789cc7906497baf8440c473d7a5a4d5214840187e02'
doc=Document(P/'MoNoTeC-2026-paper.docx')
assert len(doc.tables)==6 and len(doc.inline_shapes)==4
assert len(doc._element.xpath('.//m:oMath'))==42
assert len(doc._element.xpath('.//m:oMathPara'))==7
assert all(s._inline.docPr.get('descr') for s in doc.inline_shapes)
assert sum(p.style.name=='Image Caption' for p in doc.paragraphs)==4
assert sum(p.style.name=='Table Caption' for p in doc.paragraphs)==6
assert all(not any(token in p.text for token in ['\\begin{','\\end{','\\cite{','\\ref{']) for p in doc.paragraphs)
paper=pymupdf.open(P/'MoNoTeC-2026-paper.pdf');assert len(paper)==15
fonts={f[0] for page in paper for f in page.get_fonts(full=True)}
assert all(paper.extract_font(x)[3] for x in fonts),'Unembedded font'
response=pymupdf.open(P/'reviewer-response.pdf');assert len(response)==6
for filename in ['MoNoTeC-2026-paper.docx','reviewer-response.docx']:
    with zipfile.ZipFile(P/filename) as z:assert z.testzip() is None
response_md=(P/'reviewer-response.md').read_text(encoding='utf-8')
for reviewer,items in [(1,7),(2,4),(3,7)]:
    for n in range(1,items+1):assert f'R{reviewer}.{n}.' in response_md
assert 'verification pending' not in response_md
report={'checks':'passed','paper_pages':15,'response_pages':6,'docx_render_pages':16,
        'docx_tables':6,'docx_images':4,'omml_nodes':42,'display_equations':7,'bibliography_items':9,
        'pdf_fonts_all_embedded':True,'model_copies_unchanged':True,
        'files':{str(p.relative_to(R)).replace('\\','/'):sha(p) for p in [source,R/'monotec2026.bib',P/'MoNoTeC-2026-paper.pdf',P/'MoNoTeC-2026-paper.docx',P/'reviewer-response.md',P/'reviewer-response.pdf',P/'reviewer-response.docx']}}
(P/'checks/artifact-validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='files'}))
